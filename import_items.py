#!/usr/bin/env python3
"""Bulk import serialized inventory into Cloud Firestore.

Input workbook: stock.xlsx, sheet "Stock"
Required columns:
  SKU, PRODUCT_NAME, PRICE_VND, QUANTITY
Optional:
  ID_PREFIX  (blank => SKU)

Safety:
- Dry-run by default. Nothing is written unless --commit is passed.
- Uses a Firestore counter per ID prefix so later imports continue sequences
  without reusing ITEM_IDs.
- Writes items in batches (default 400 writes per batch).
- Exports a CSV containing every generated ITEM_ID + QR_CONTENT.\n- Creates updatedAt on products/items so Google Sheets incremental sync can detect changes.

Example:
  python import_items.py stock.xlsx --service-account serviceAccountKey.json
  python import_items.py stock.xlsx --service-account serviceAccountKey.json --commit

You may omit --service-account if Application Default Credentials are already set.
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import sys
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Iterable, List, Optional

import firebase_admin
from firebase_admin import credentials
from firebase_admin import firestore as admin_firestore
from google.cloud import firestore as gcfirestore
from openpyxl import load_workbook

DEFAULT_PROJECT_ID = "item-level-tracking"
DEFAULT_BATCH_SIZE = 400
ID_WIDTH = 4
VALID_ID_RE = re.compile(r"^[A-Z0-9][A-Z0-9_-]*$")


@dataclass
class StockRow:
    sku: str
    product_name: str
    price_vnd: int
    quantity: int
    id_prefix: str


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Generate serialized ITEM_IDs and bulk-import them into Firestore."
    )
    p.add_argument("input", help="Path to stock.xlsx")
    p.add_argument(
        "--sheet",
        default="Stock",
        help='Worksheet name (default: "Stock")',
    )
    p.add_argument(
        "--project-id",
        default=DEFAULT_PROJECT_ID,
        help=f"Firebase / GCP project ID (default: {DEFAULT_PROJECT_ID})",
    )
    p.add_argument(
        "--service-account",
        help="Path to service-account JSON. If omitted, Application Default Credentials are used.",
    )
    p.add_argument(
        "--batch-size",
        type=int,
        default=DEFAULT_BATCH_SIZE,
        help=f"Firestore writes per commit batch (default: {DEFAULT_BATCH_SIZE}; keep <= 450).",
    )
    p.add_argument(
        "--output",
        help="Generated item CSV path. Default: generated_items_<timestamp>.csv",
    )
    p.add_argument(
        "--commit",
        action="store_true",
        help="Actually write to Firestore. Without this flag the script is a dry run.",
    )
    return p.parse_args()


def normalize_text(value) -> str:
    return str(value or "").strip()


def normalize_code(value, field_name: str) -> str:
    code = normalize_text(value).upper()
    if not code:
        raise ValueError(f"{field_name} cannot be blank")
    if not VALID_ID_RE.fullmatch(code):
        raise ValueError(
            f"{field_name}={code!r} is invalid. Use only A-Z, 0-9, hyphen (-), underscore (_)."
        )
    if "/" in code:
        raise ValueError(f"{field_name} cannot contain '/'")
    return code


def as_positive_int(value, field_name: str) -> int:
    if value is None or value == "":
        raise ValueError(f"{field_name} cannot be blank")
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise ValueError(f"{field_name} must be an integer, got {value!r}")
    if number <= 0:
        raise ValueError(f"{field_name} must be > 0, got {number}")
    return number


def as_nonnegative_int(value, field_name: str) -> int:
    if value is None or value == "":
        raise ValueError(f"{field_name} cannot be blank")
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise ValueError(f"{field_name} must be an integer, got {value!r}")
    if number < 0:
        raise ValueError(f"{field_name} must be >= 0, got {number}")
    return number


def load_stock_rows(path: Path, sheet_name: str) -> List[StockRow]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Worksheet {sheet_name!r} not found. Available: {', '.join(wb.sheetnames)}"
        )
    ws = wb[sheet_name]

    rows = ws.iter_rows(values_only=True)
    try:
        headers = [normalize_text(x).upper() for x in next(rows)]
    except StopIteration:
        raise ValueError("Workbook is empty")

    required = ["SKU", "PRODUCT_NAME", "PRICE_VND", "QUANTITY"]
    missing = [h for h in required if h not in headers]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")

    idx = {name: headers.index(name) for name in headers if name}
    parsed: List[StockRow] = []

    for excel_row_num, row in enumerate(rows, start=2):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue

        try:
            sku = normalize_code(row[idx["SKU"]], "SKU")
            product_name = normalize_text(row[idx["PRODUCT_NAME"]])
            if not product_name:
                raise ValueError("PRODUCT_NAME cannot be blank")
            price = as_nonnegative_int(row[idx["PRICE_VND"]], "PRICE_VND")
            quantity = as_positive_int(row[idx["QUANTITY"]], "QUANTITY")
            raw_prefix = row[idx["ID_PREFIX"]] if "ID_PREFIX" in idx else None
            prefix = normalize_code(raw_prefix, "ID_PREFIX") if normalize_text(raw_prefix) else sku
        except Exception as exc:
            raise ValueError(f"Row {excel_row_num}: {exc}") from exc

        parsed.append(
            StockRow(
                sku=sku,
                product_name=product_name,
                price_vnd=price,
                quantity=quantity,
                id_prefix=prefix,
            )
        )

    if not parsed:
        raise ValueError("No stock rows found")

    return consolidate_rows(parsed)


def consolidate_rows(rows: List[StockRow]) -> List[StockRow]:
    """Merge repeated SKUs while guarding against conflicting metadata."""
    by_sku: Dict[str, StockRow] = {}
    prefix_owner: Dict[str, str] = {}

    for r in rows:
        owner = prefix_owner.get(r.id_prefix)
        if owner and owner != r.sku:
            raise ValueError(
                f"ID_PREFIX {r.id_prefix!r} is shared by SKU {owner!r} and {r.sku!r}. "
                "Each SKU must have a unique ID_PREFIX."
            )
        prefix_owner[r.id_prefix] = r.sku

        existing = by_sku.get(r.sku)
        if existing is None:
            by_sku[r.sku] = StockRow(**r.__dict__)
            continue

        if (
            existing.product_name != r.product_name
            or existing.price_vnd != r.price_vnd
            or existing.id_prefix != r.id_prefix
        ):
            raise ValueError(
                f"Duplicate SKU {r.sku!r} has conflicting PRODUCT_NAME / PRICE_VND / ID_PREFIX"
            )
        existing.quantity += r.quantity

    return list(by_sku.values())


def init_firestore(project_id: str, service_account: Optional[str]):
    if firebase_admin._apps:
        return admin_firestore.client()

    if service_account:
        cred = credentials.Certificate(service_account)
    else:
        cred = credentials.ApplicationDefault()

    firebase_admin.initialize_app(cred, {"projectId": project_id})
    return admin_firestore.client()


def reserve_sequence(db, prefix: str, quantity: int) -> int:
    """Atomically reserve [start, start+quantity-1] from counters/{prefix}."""
    counter_ref = db.collection("counters").document(prefix)
    transaction = db.transaction()

    @gcfirestore.transactional
    def _reserve(txn):
        snap = counter_ref.get(transaction=txn)
        start = 1
        if snap.exists:
            data = snap.to_dict() or {}
            start = int(data.get("nextNumber", 1))

        txn.set(
            counter_ref,
            {
                "nextNumber": start + quantity,
                "updatedAt": gcfirestore.SERVER_TIMESTAMP,
            },
            merge=True,
        )
        return start

    return _reserve(transaction)


def chunked(items: List[dict], size: int) -> Iterable[List[dict]]:
    for i in range(0, len(items), size):
        yield items[i : i + size]


def generate_preview(rows: List[StockRow]) -> List[dict]:
    """Dry-run IDs start at 1 per prefix because no counter is modified."""
    output: List[dict] = []
    for r in rows:
        for seq in range(1, r.quantity + 1):
            item_id = f"{r.id_prefix}-{seq:0{ID_WIDTH}d}"
            output.append(
                {
                    "ITEM_ID": item_id,
                    "SKU": r.sku,
                    "PRODUCT_NAME": r.product_name,
                    "PRICE_VND": r.price_vnd,
                    "STATUS": "IN_STOCK",
                    "QR_CONTENT": item_id,
                    "BATCH_ID": "DRY-RUN",
                }
            )
    return output


def commit_import(db, rows: List[StockRow], batch_size: int) -> List[dict]:
    if batch_size < 1 or batch_size > 450:
        raise ValueError("--batch-size must be between 1 and 450")

    batch_id = datetime.now(timezone.utc).strftime("BATCH-%Y%m%dT%H%M%SZ-") + uuid.uuid4().hex[:8]
    generated: List[dict] = []

    # Upsert product documents first.
    product_batch = db.batch()
    for r in rows:
        product_ref = db.collection("products").document(r.sku)
        product_batch.set(
            product_ref,
            {
                "sku": r.sku,
                "name": r.product_name,
                "price": r.price_vnd,
                "updatedAt": gcfirestore.SERVER_TIMESTAMP,
            },
            merge=True,
        )
    product_batch.commit()

    # Reserve ranges and prepare item documents.
    item_payloads: List[dict] = []
    for r in rows:
        start = reserve_sequence(db, r.id_prefix, r.quantity)
        end = start + r.quantity - 1
        print(f"Reserved {r.id_prefix}: {start:0{ID_WIDTH}d}..{end:0{ID_WIDTH}d}")

        for seq in range(start, end + 1):
            item_id = f"{r.id_prefix}-{seq:0{ID_WIDTH}d}"
            item_payloads.append(
                {
                    "ITEM_ID": item_id,
                    "SKU": r.sku,
                    "PRODUCT_NAME": r.product_name,
                    "PRICE_VND": r.price_vnd,
                    "STATUS": "IN_STOCK",
                    "QR_CONTENT": item_id,
                    "BATCH_ID": batch_id,
                }
            )

    # Commit serialized item docs in safe-sized batches.
    total = len(item_payloads)
    written = 0
    for group in chunked(item_payloads, batch_size):
        batch = db.batch()
        for x in group:
            item_ref = db.collection("items").document(x["ITEM_ID"])
            batch.set(
                item_ref,
                {
                    "itemId": x["ITEM_ID"],
                    "sku": x["SKU"],
                    "status": "IN_STOCK",
                    "orderId": None,
                    "soldAt": None,
                    "soldBy": None,
                    "soldDeviceId": None,
                    "updatedAt": gcfirestore.SERVER_TIMESTAMP,
                    "restockedAt": None,
                    "restockedBy": None,
                    "restockedDeviceId": None,
                    "restockReason": None,
                    "lastOrderId": None,
                    "lastSoldAt": None,
                    "lastSoldBy": None,
                    "lastSoldDeviceId": None,
                    "batchId": batch_id,
                    "createdAt": gcfirestore.SERVER_TIMESTAMP,
                },
            )
        batch.commit()
        written += len(group)
        print(f"Committed {written}/{total} item documents")

    generated.extend(item_payloads)
    return generated


def write_csv(records: List[dict], output_path: Path) -> None:
    fields = [
        "ITEM_ID",
        "SKU",
        "PRODUCT_NAME",
        "PRICE_VND",
        "STATUS",
        "QR_CONTENT",
        "BATCH_ID",
    ]
    with output_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(records)


def main() -> int:
    args = parse_args()
    input_path = Path(args.input)
    if not input_path.exists():
        print(f"ERROR: input file not found: {input_path}", file=sys.stderr)
        return 2

    try:
        rows = load_stock_rows(input_path, args.sheet)
    except Exception as exc:
        print(f"ERROR reading workbook: {exc}", file=sys.stderr)
        return 2

    total_qty = sum(r.quantity for r in rows)
    print("\nValidated stock:")
    for r in rows:
        print(
            f"  {r.sku:<20} qty={r.quantity:<6} price={r.price_vnd:<10} prefix={r.id_prefix}"
        )
    print(f"Total physical items: {total_qty}\n")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = Path(args.output or f"generated_items_{timestamp}.csv")

    if not args.commit:
        print("DRY RUN: no Firestore data will be changed.")
        records = generate_preview(rows)
        write_csv(records, output_path)
        print(f"Preview CSV written: {output_path.resolve()}")
        print("\nWhen ready, rerun the same command with --commit.")
        return 0

    try:
        db = init_firestore(args.project_id, args.service_account)
        records = commit_import(db, rows, args.batch_size)
        write_csv(records, output_path)
    except Exception as exc:
        print(f"ERROR during Firestore import: {exc}", file=sys.stderr)
        print(
            "IMPORTANT: if the failure happened after sequence reservation, do not manually rewind counters. "
            "Gaps in ITEM_IDs are safe; duplicate IDs are not.",
            file=sys.stderr,
        )
        return 1

    print("\nIMPORT COMPLETE")
    print(f"Firestore items created: {len(records)}")
    print(f"Generated CSV: {output_path.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
